#!/usr/bin/env python3
"""Convert the captured AnswerSocrates export into an attested PAA artifact."""

from __future__ import annotations

import csv
import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))

from data_sources.modules.paa_provenance_guard import (
    build_answersocrates_artifact,
    write_answersocrates_artifact,
)


OUTPUT_DIR = REPO / "research" / "best-plumbing-job-management-software-remediation-2026-09-03"
SOURCE_XLSX = OUTPUT_DIR / "raw" / "answersocrates" / "AnswerSocrates-US-en-plumbing-job-management-software-2026-09-03.xlsx"
ARTIFACT_PATH = OUTPUT_DIR / "answersocrates-paa.json"
CSV_PATH = OUTPUT_DIR / "answersocrates-paa-questions.csv"
SOURCE_EVIDENCE_PATH = OUTPUT_DIR / "answersocrates-source-evidence.json"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_paa_rows() -> list[dict[str, str]]:
    workbook = load_workbook(SOURCE_XLSX, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    records = []
    for values in rows:
        record = {
            header: str(value).strip() if value is not None else ""
            for header, value in zip(headers, values)
        }
        if record.get("Type", "").casefold() == "paa":
            records.append(record)
    return records


def main() -> int:
    rows = read_paa_rows()
    questions = [row["Question"] for row in rows if row.get("Question")]
    if len(questions) != 4:
        raise RuntimeError(f"Expected four PAA questions in the captured export, found {len(questions)}")

    artifact = build_answersocrates_artifact(
        query="plumbing job management software",
        collection_date="2026-09-03",
        eligible_questions=questions,
        ineligible_fragments=(),
        run_id="plumbing-serp-recovery-answersocrates-2026-09-03",
        started_at="2026-09-03T18:45:18Z",
        completed_at="2026-09-03T18:46:05Z",
        tool={"name": "playwright_cli", "version": "0.1.19"},
    )
    source_export = {
        "path": str(SOURCE_XLSX.relative_to(REPO)).replace("\\", "/"),
        "sha256": sha256(SOURCE_XLSX),
        "format": "xlsx",
        "country": "United States",
        "language": "English",
    }
    write_answersocrates_artifact(ARTIFACT_PATH, artifact)
    SOURCE_EVIDENCE_PATH.write_text(
        json.dumps(
            {
                "schema": "simpro-answersocrates-source-evidence/v1",
                "query": "plumbing job management software",
                "collection_date": "2026-09-03",
                "collection_runtime": {"name": "playwright_cli", "version": "0.1.19"},
                "started_at": "2026-09-03T18:45:18Z",
                "completed_at": "2026-09-03T18:46:05Z",
                "source_export": source_export,
                "screenshot": "research/best-plumbing-job-management-software-remediation-2026-09-03/raw/answersocrates/plumbing-job-management-software-2026-09-03.png",
                "paa_question_count": len(questions),
                "artifact": str(ARTIFACT_PATH.relative_to(REPO)).replace("\\", "/"),
                "artifact_receipt_hash": artifact["run_receipt"]["receipt_hash"],
            },
            indent=2,
            ensure_ascii=True,
        )
        + "\n",
        encoding="utf-8",
    )

    with CSV_PATH.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["source_type", "modifier", "question", "keyword_label"],
        )
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    "source_type": row.get("Type", ""),
                    "modifier": row.get("Modifier", ""),
                    "question": row.get("Question", ""),
                    "keyword_label": row.get("Keyword Label", ""),
                }
            )

    print(
        json.dumps(
            {
                "artifact": str(ARTIFACT_PATH),
                "questions": len(questions),
                "receipt_hash": artifact["run_receipt"]["receipt_hash"],
                "source_export_sha256": source_export["sha256"],
            }
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
